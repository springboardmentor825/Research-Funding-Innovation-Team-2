"""Streaming readers for the local, read-only public datasets.

Dataset rows intentionally never enter MongoDB or a process-wide data cache.
"""

from __future__ import annotations

import csv
import math
import os
from datetime import date, datetime
from pathlib import Path
from threading import Lock
from typing import Any, Iterator

from openpyxl import load_workbook


APP_DIR = Path(__file__).resolve().parents[1]
DEFAULT_DATA_DIR = APP_DIR / "data"
DATA_DIR = Path(os.getenv("DATA_DIR", str(DEFAULT_DATA_DIR))).expanduser()

# The legacy file is misspelled in the supplied repository.  Prefer the documented
# spelling when it is added later, while retaining compatibility with this checkout.
DATASETS: dict[str, dict[str, Any]] = {
    "datasets": {"files": ("datasets.csv", "datsets.csv"), "type": "csv"},
    "grants": {"files": ("grants.csv",), "type": "csv"},
    "patents": {"files": ("patentbrief-patents.csv",), "type": "csv"},
    "technology": {"files": ("technology_dataset.xlsx",), "type": "excel"},
}

_count_cache: dict[str, tuple[int, int, int]] = {}
_cache_lock = Lock()


class DatasetError(Exception):
    """Base error for a fixed-registry dataset request."""


class UnknownDatasetError(DatasetError):
    pass


class DatasetFileMissingError(DatasetError):
    pass


class InvalidDatasetFilterError(DatasetError):
    pass


def get_dataset_config(dataset: str) -> dict[str, Any]:
    config = DATASETS.get(dataset)
    if config is None:
        raise UnknownDatasetError(f"Unknown dataset: {dataset}")
    return config


def get_dataset_path(dataset: str) -> Path:
    config = get_dataset_config(dataset)
    for filename in config["files"]:
        candidate = DATA_DIR / filename
        if candidate.is_file():
            return candidate
    expected = DATA_DIR / config["files"][0]
    raise DatasetFileMissingError(f"Dataset file is missing: {expected.name}")


def _clean_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, str):
        value = value.strip()
        return value if value else None
    # openpyxl may return scalar types which FastAPI can serialize, but converting
    # unusual scalar objects here keeps this reader independent of those details.
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _csv_rows(path: Path) -> Iterator[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as source:
        reader = csv.DictReader(source)
        for row in reader:
            # A malformed CSV row can create a None key for overflow fields.
            yield {str(key): _clean_value(value) for key, value in row.items() if key is not None}


def _excel_headers(path: Path) -> list[str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        return [str(value).strip() for value in first_row if value is not None]
    finally:
        workbook.close()


def _excel_rows(path: Path) -> Iterator[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        raw_headers = next(rows, ())
        headers = [str(value).strip() if value is not None else f"column_{index + 1}"
                   for index, value in enumerate(raw_headers)]
        for values in rows:
            yield {headers[index]: _clean_value(value) for index, value in enumerate(values)}
    finally:
        workbook.close()


def iter_rows(dataset: str) -> Iterator[dict[str, Any]]:
    config = get_dataset_config(dataset)
    path = get_dataset_path(dataset)
    return _csv_rows(path) if config["type"] == "csv" else _excel_rows(path)


def get_columns(dataset: str) -> list[str]:
    config = get_dataset_config(dataset)
    path = get_dataset_path(dataset)
    if config["type"] == "excel":
        return _excel_headers(path)
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as source:
        return [column.strip() for column in (csv.DictReader(source).fieldnames or []) if column]


def count_rows(dataset: str) -> int:
    """Return a cached count, invalidated whenever file size or mtime changes."""
    path = get_dataset_path(dataset)
    stat = path.stat()
    fingerprint = (stat.st_mtime_ns, stat.st_size)
    cache_key = str(path.resolve())
    with _cache_lock:
        cached = _count_cache.get(cache_key)
        if cached and cached[:2] == fingerprint:
            return cached[2]

    config = get_dataset_config(dataset)
    if config["type"] == "csv":
        # DictReader handles quoted newlines correctly, unlike counting text lines.
        count = sum(1 for _ in _csv_rows(path))
    else:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook[workbook.sheetnames[0]]
            count = max(sheet.max_row - 1, 0)
        finally:
            workbook.close()

    with _cache_lock:
        _count_cache[cache_key] = (*fingerprint, count)
    return count


def validate_filters(dataset: str, filters: dict[str, str]) -> dict[str, str]:
    columns = get_columns(dataset)
    by_lower_name = {column.casefold(): column for column in columns}
    invalid = [name for name in filters if name.casefold() not in by_lower_name]
    if invalid:
        available = ", ".join(columns)
        raise InvalidDatasetFilterError(
            f"Unsupported filter(s): {', '.join(invalid)}. Available columns: {available}"
        )
    return {by_lower_name[name.casefold()]: value for name, value in filters.items()}


def _matches(row: dict[str, Any], query: str | None, filters: dict[str, str]) -> bool:
    for column, wanted in filters.items():
        value = row.get(column)
        if value is None or str(value).strip().casefold() != wanted.strip().casefold():
            return False
    if not query:
        return True
    needle = query.strip().casefold()
    if not needle:
        return True
    # Avoid identifiers and URLs, which are not meaningful free-text search fields.
    searchable = (
        value for column, value in row.items()
        if not column.casefold().endswith(("id", "number", "url"))
        and column.casefold() not in {"references", "n_citation", "times_cited"}
    )
    return any(value is not None and needle in str(value).casefold() for value in searchable)


def search_rows(
    dataset: str, query: str | None, filters: dict[str, str], page: int, limit: int
) -> tuple[int, list[dict[str, Any]]]:
    """Scan once, retaining only the requested page while counting all matches."""
    offset = (page - 1) * limit
    total = 0
    result: list[dict[str, Any]] = []
    for row in iter_rows(dataset):
        if not _matches(row, query, filters):
            continue
        if offset <= total < offset + limit:
            result.append(row)
        total += 1
    return total, result


def paginated_rows(dataset: str, page: int, limit: int) -> tuple[int, list[dict[str, Any]]]:
    total = count_rows(dataset)
    offset = (page - 1) * limit
    result: list[dict[str, Any]] = []
    for index, row in enumerate(iter_rows(dataset)):
        if index < offset:
            continue
        if len(result) == limit:
            break
        result.append(row)
    return total, result
